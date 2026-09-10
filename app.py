import streamlit as st
import pandas as pd
import datetime
import calendar
import random
import sqlite3
from io import BytesIO
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

st.set_page_config(page_title="Sistem Penugasan SBCC IPK Selangor", layout="wide")

# ==========================================
# 1. FUNGSI PANGKALAN DATA (SQLITE)
# ==========================================
DB_FILE = "jadual_penugasan.db"

def init_db():
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    # Ditambah lajur seksyen_pegawai dan seksyen_penolong
    c.execute('''
        CREATE TABLE IF NOT EXISTS rekod_jadual (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            bulan INTEGER,
            tahun INTEGER,
            tarikh TEXT,
            hari TEXT,
            syif TEXT,
            nama_pegawai TEXT,
            seksyen_pegawai TEXT,
            no_tel_pegawai TEXT,
            nama_penolong TEXT,
            seksyen_penolong TEXT,
            no_tel_penolong TEXT,
            tarikh_disahkan TEXT
        )
    ''')
    conn.commit()
    conn.close()

def simpan_jadual_db(bulan, tahun, df_jadual):
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("DELETE FROM rekod_jadual WHERE bulan = ? AND tahun = ?", (bulan, tahun))
    
    tarikh_sekarang = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for _, row in df_jadual.iterrows():
        c.execute('''
            INSERT INTO rekod_jadual (bulan, tahun, tarikh, hari, syif, nama_pegawai, seksyen_pegawai, no_tel_pegawai, nama_penolong, seksyen_penolong, no_tel_penolong, tarikh_disahkan)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (bulan, tahun, row['Tarikh'], row['Hari'], row['Syif / Masa'], 
              row['Nama Pegawai'], row['Seksyen Pegawai'], row['No. Tel Pegawai'], 
              row['Nama Penolong'], row['Seksyen Penolong'], row['No. Tel Penolong'], tarikh_sekarang))
    
    conn.commit()
    conn.close()

def padam_jadual_db(bulan, tahun):
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("DELETE FROM rekod_jadual WHERE bulan = ? AND tahun = ?", (bulan, tahun))
    conn.commit()
    conn.close()

def ambil_jadual_db(bulan, tahun):
    conn = sqlite3.connect(DB_FILE)
    query = """
        SELECT tarikh AS Tarikh, hari AS Hari, syif AS 'Syif / Masa', 
               nama_pegawai AS 'Nama Pegawai', seksyen_pegawai AS 'Seksyen Pegawai', no_tel_pegawai AS 'No. Tel Pegawai',
               nama_penolong AS 'Nama Penolong', seksyen_penolong AS 'Seksyen Penolong', no_tel_penolong AS 'No. Tel Penolong',
               tarikh_disahkan AS 'Tarikh Disahkan' 
        FROM rekod_jadual WHERE bulan = ? AND tahun = ?
    """
    df = pd.read_sql_query(query, conn, params=(bulan, tahun))
    conn.close()
    return df

def is_pegawai_wanita(nama):
    nama_upper = nama.upper()
    kata_kunci_wanita = [" BINTI ", " BT ", " A/P ", " KAUR ", " LO YEE TYNG ", " YVONNE KUA "]
    return any(k in f" {nama_upper} " for k in kata_kunci_wanita)

# ----------------------------------------------------
# FUNGSI LAMPIRAN 'A' (HTML & WORD)
# ----------------------------------------------------
def jana_html_lampiran_a(df_jadual, bulan, tahun):
    nama_bulan_map = {1: "JAN", 2: "FEB", 3: "MAC", 4: "APR", 5: "MEI", 6: "JUN", 7: "JUL", 8: "OGOS", 9: "SEP", 10: "OKT", 11: "NOV", 12: "DIS"}
    bln_str = nama_bulan_map.get(bulan, "SEP")
    
    html = f"""
    <!DOCTYPE html>
    <html>
    <head>
    <style>
        @page {{ size: A4 portrait; margin: 10mm; }}
        body {{ font-family: Arial, sans-serif; font-size: 10px; color: #000; }}
        .header {{ text-align: center; font-weight: bold; margin-bottom: 12px; font-size: 11px; line-height: 1.3; }}
        .lampiran {{ float: right; font-weight: bold; text-decoration: underline; font-size: 10px; }}
        table {{ width: 100%; border-collapse: collapse; margin-top: 5px; page-break-inside: auto; }}
        tr {{ page-break-inside: avoid; page-break-after: auto; }}
        th, td {{ border: 1px solid black; padding: 5px 6px; text-align: center; vertical-align: middle; }}
        th {{ background-color: #e0e0e0; font-weight: bold; font-size: 10px; }}
        .text-left {{ text-align: left; }}
        .jawatan {{ font-size: 9px; color: #333; font-style: italic; }}
        .seksyen {{ font-weight: bold; color: #b30000; }}
    </style>
    </head>
    <body>
        <div class="lampiran">LAMPIRAN 'A'</div>
        <div style="clear:both;"></div>
        <div class="header">
            JADUAL PEGAWAI & PENOLONG PEGAWAI BERTUGAS<br>
            SBCC IPK SELANGOR BAGI BULAN {bln_str} {tahun}
        </div>
        <table>
            <thead>
                <tr>
                    <th style="width: 5%;">BIL</th>
                    <th style="width: 18%;">TKH / MASA</th>
                    <th style="width: 12%;">HARI</th>
                    <th style="width: 65%;">PEGAWAI & PENOLONG PEGAWAI BERTUGAS</th>
                </tr>
            </thead>
            <tbody>
    """
    
    for idx, row in df_jadual.iterrows():
        bil = f"{idx+1:02d}."
        masa = "1700 - 0800"
        if "Siang" in row['Syif / Masa']: masa = "0800 - 2000"
        elif "Malam" in row['Syif / Masa']: masa = "2000 - 0800"
            
        tkh_masa = f"{row['Tarikh']}<br>{masa}"
        hari = row['Hari'].upper()
        
        petak_gabung = f"""
            <b>1. {row['Nama Pegawai']}</b> <span class="seksyen">({row['Seksyen Pegawai']})</span> <span class="jawatan">(Pegawai Bertugas)</span><br>
            &nbsp;&nbsp;&nbsp;&nbsp;<b>No. Tel:</b> {row['No. Tel Pegawai']}<br>
            <div style="margin-top:3px; border-top: 1px dotted #ccc; padding-top: 3px;">
            <b>2. {row['Nama Penolong']}</b> <span class="seksyen">({row['Seksyen Penolong']})</span> <span class="jawatan">(Penolong Pegawai Bertugas)</span><br>
            &nbsp;&nbsp;&nbsp;&nbsp;<b>No. Tel:</b> {row['No. Tel Penolong']}
            </div>
        """
        
        html += f"""
            <tr>
                <td>{bil}</td>
                <td>{tkh_masa}</td>
                <td>{hari}</td>
                <td class="text-left">{petak_gabung}</td>
            </tr>
        """
        
    html += """
            </tbody>
        </table>
    </body>
    </html>
    """
    return html

def jana_word_lampiran_a(df_jadual, bulan, tahun):
    nama_bulan_map = {1: "JAN", 2: "FEB", 3: "MAC", 4: "APR", 5: "MEI", 6: "JUN", 7: "JUL", 8: "OGOS", 9: "SEP", 10: "OKT", 11: "NOV", 12: "DIS"}
    bln_str = nama_bulan_map.get(bulan, "SEP")
    
    doc = Document()
    
    # Tetapan Margin A4
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(1.27)
        section.bottom_margin = Cm(1.27)
        section.left_margin = Cm(1.27)
        section.right_margin = Cm(1.27)

    # Header Lampiran
    p_lamp = doc.add_paragraph()
    p_lamp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_lamp = p_lamp.add_run("LAMPIRAN 'A'")
    run_lamp.bold = True
    run_lamp.underline = True
    
    # Header Tajuk
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_head = p_head.add_run(f"JADUAL PEGAWAI & PENOLONG PEGAWAI BERTUGAS\nSBCC IPK SELANGOR BAGI BULAN {bln_str} {tahun}")
    run_head.bold = True
    
    # Cipta Jadual
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'BIL'
    hdr_cells[1].text = 'TKH / MASA'
    hdr_cells[2].text = 'HARI'
    hdr_cells[3].text = 'PEGAWAI & PENOLONG PEGAWAI BERTUGAS'
    
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True

    for idx, row in df_jadual.iterrows():
        row_cells = table.add_row().cells
        row_cells[0].text = f"{idx+1:02d}."
        
        masa = "1700 - 0800"
        if "Siang" in row['Syif / Masa']: masa = "0800 - 2000"
        elif "Malam" in row['Syif / Masa']: masa = "2000 - 0800"
            
        row_cells[1].text = f"{row['Tarikh']}\n{masa}"
        row_cells[2].text = row['Hari'].upper()
        
        pegawai_txt = f"1. {row['Nama Pegawai']} ({row['Seksyen Pegawai']})\n    (Pegawai Bertugas)\n    No. Tel: {row['No. Tel Pegawai']}\n"
        penolong_txt = f"2. {row['Nama Penolong']} ({row['Seksyen Penolong']})\n    (Penolong Pegawai Bertugas)\n    No. Tel: {row['No. Tel Penolong']}"
        
        row_cells[3].text = pegawai_txt + penolong_txt
        
        # Center align untuk 3 lajur pertama
        for i in range(3):
            for p in row_cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
    # Simpan ke dalam buffer memori
    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()

init_db()

# ==========================================
# 2. SENARAI DATA PEGAWAI & PENOLONG PEGAWAI
# Sila ubah nilai "seksyen" mengikut cawangan anggota masing-masing (E1/E2/E3 dll)
# ==========================================
st.title("📋 Sistem Penugasan SBCC IPK Selangor")

PEGAWAI = [
    {"nama": "INSP ADDLYZAN BIN ABD MANAP", "seksyen": "E3", "no_tel": "012-4739701"},
    {"nama": "INSP FAIZ BIN BASAR", "seksyen": "E4", "no_tel": "017-9457941"},
    {"nama": "INSP MOHD SHAHRILBUN BIN AMDAN", "seksyen": "E3", "no_tel": "017-5800578"},
    {"nama": "INSP FARAH NADJWA BT MOHD LASA", "seksyen": "E5", "no_tel": "013-9195659"},
    {"nama": "INSP SABARI BIN BUJANG", "seksyen": "E6", "no_tel": "016-7135975"},
    {"nama": "INSP KASHMINDEJRIT KAUR A/P CHARLES", "seksyen": "E4", "no_tel": "016-6738465"},
    {"nama": "INSP NORAINUN MUBIN BT ANUAR", "seksyen": "E3", "no_tel": "012-7493780"},
    {"nama": "INSP MUHAMAD IQRAM BIN MOHAMAD BUKHARI", "seksyen": "E3", "no_tel": "013-7044459"},
    {"nama": "INSP MOHAMAD IQBAL IBRAHIM BIN DAUD", "seksyen": "E4", "no_tel": "012-6467775"},
    {"nama": "INSP AMI RUSLAN BIN MOHD RUSLAN", "seksyen": "E6", "no_tel": "017-9501506"},
    {"nama": "INSP MOHD RAZIF BIN ROSSMAN", "seksyen": "E5", "no_tel": "017-9196166"},
    {"nama": "INSP AZFIZI BIN AZIZ", "seksyen": "E3", "no_tel": "014-8340944"},
    {"nama": "INSP MUHAMAD ARIFF BIN ABD RAHIM", "seksyen": "E4", "no_tel": "010-4017822"},
    {"nama": "INSP MOHD RUSDI BIN HASSAN", "seksyen": "E3", "no_tel": "010-4535060"},
    {"nama": "INSP ABDUL MUIN BIN AB AZIZ", "seksyen": "E3", "no_tel": "019-4494500"},
    {"nama": "INSP MOHD SHAHRIL BIN MOHAMED RESALI", "seksyen": "E4", "no_tel": "012-9096486"},
    {"nama": "INSP NOR HAIZAH BINTI YACOB", "seksyen": "E6", "no_tel": "012-9979790"},
    {"nama": "INSP KUKENDRAN A/L YOGENDRAN", "seksyen": "E3", "no_tel": "012-9062264"},
    {"nama": "INSP MUHAMMAD NOR QHAIREIL BIN NOR AZMI", "seksyen": "E5", "no_tel": "011-16554250"},
    {"nama": "INSP KHAIRUL HAFIZ BIN SHAFFIE AHMED", "seksyen": "E3", "no_tel": "017-4308434"},
    {"nama": "INSP EZYANTI BT MUHAMAD", "seksyen": "E4", "no_tel": "019-9241360"},
    {"nama": "INSP MOHAMAD SYARUL FADLI BIN MAT HUSAIN", "seksyen": "E6", "no_tel": "011-10042323"},
    {"nama": "INSP MUHAMMAD AFIQ BIN NORDIN", "seksyen": "E3", "no_tel": "010-4641637"},
    {"nama": "INSP MOHD JEFFERE BIN ALI", "seksyen": "E3", "no_tel": "014-2348186"},
    {"nama": "INSP NUR AADILA BINTI ABDUL MALEK", "seksyen": "E4", "no_tel": "014-3383604"},
    {"nama": "INSP SHURENDREN A/L JAYARAMAN", "seksyen": "E5", "no_tel": "019-2823312"},
    {"nama": "INSP MOHD FAIZ BIN ESA", "seksyen": "E3", "no_tel": "010-7978603"},
    {"nama": "INSP ROHAIZASHAFIKA BINTI MOHAMAD RADZALI", "seksyen": "E6", "no_tel": "016-2448782"},
    {"nama": "INSP DZUL FADHLI BIN ABDUL HALIM", "seksyen": "E3", "no_tel": "014-6689803"},
    {"nama": "INSP NORHAZIERAM BINTI ZAKARIA", "seksyen": "E4", "no_tel": "014-9207312"},
    {"nama": "INSP WAN MUHAMAD MUSTAQIM BIN ABDUL RAHMAN", "seksyen": "E3", "no_tel": "012-9212972"}
]

PENOLONG_PEGAWAI = [
    {"nama": "SJN/D HAFIZATUL AMIRA BT ADBUL MANAF", "seksyen": "E3", "no_tel": "019-5493007"},
    {"nama": "SJN/D AHMAD ZAMANE BIN OMAR", "seksyen": "E4", "no_tel": "013-4939452"},
    {"nama": "SJN/D ZAMRI BIN BAKARI", "seksyen": "E6", "no_tel": "013-2893878"},
    {"nama": "KPL/D MUHAMMAD AZKA FATIN BIN MOHD NORDIN", "seksyen": "E3", "no_tel": "016-2198068"},
    {"nama": "KPL/D MUHAMMAD ZULFFY BIN RASHID", "seksyen": "E5", "no_tel": "013-4080883"},
    {"nama": "KPL/D MUHAMMAD FIRDAUS BIN AMRAN", "seksyen": "E3", "no_tel": "013-7783677"},
    {"nama": "KPL/D HAZIZAT BIN ATTAN", "seksyen": "E4", "no_tel": "013-5744426"},
    {"nama": "SJN/D ONG WEI TAU", "seksyen": "E3", "no_tel": "012-8121981"},
    {"nama": "KPL/D MOHD FAIZAL BIN TIMAN", "seksyen": "E6", "no_tel": "017-7455348"},
    {"nama": "KPL/D JEEVAN A/L VADIAHLAGAN", "seksyen": "E3", "no_tel": "013-5638303"},
    {"nama": "KPL/D MOHD SYAIFUL AZHAR BIN MOHD HANAFIAH", "seksyen": "E4", "no_tel": "011-51105497"},
    {"nama": "SJN/D SYEAFUUL AZNRY BIN ZAMBRI", "seksyen": "E5", "no_tel": "012-7637889"},
    {"nama": "SJN/D ZAMZUHAIRI BIN TOMIRAN", "seksyen": "E3", "no_tel": "012-2712700"},
    {"nama": "SJN/D YUSRI BIN WAHAB", "seksyen": "E4", "no_tel": "019-3876720"},
    {"nama": "KPL/D MOHD YUSRI BIN ABU SEMAN", "seksyen": "E6", "no_tel": "017-9278564"},
    {"nama": "KPL/D MUHAMMAD AZRUL SHAFIQ BIN ISMAN", "seksyen": "E3", "no_tel": "019-5898658"},
    {"nama": "KPL/D FARIEZUL NAIM BIN MOHD SANI", "seksyen": "E5", "no_tel": "013-6399275"},
    {"nama": "L/KPL/D MOHAMMAD ARRYAN JUING BIN ABDULLAH", "seksyen": "E3", "no_tel": "019-4334734"},
    {"nama": "SJN/D HAFENDI BIN HANAFIAH", "seksyen": "E4", "no_tel": "012-3784305"},
    {"nama": "SJN/D EFFARINA BIN AHMAD SUHAIMI", "seksyen": "E6", "no_tel": "018-9154504"},
    {"nama": "KPL/D NOOR HAZWANI BINTI MUHAMAD RUSLI", "seksyen": "E3", "no_tel": "019-2217116"},
    {"nama": "KPL/D WELLINGTON RICHARD ANAK SADOK", "seksyen": "E4", "no_tel": "011-29994591"},
    {"nama": "SJN/D NOOR HAKIMI BIN ZULKIFLI", "seksyen": "E5", "no_tel": "012-6949107"},
    {"nama": "SJN/D SUYANI BINTI MAT SAAD", "seksyen": "E3", "no_tel": "012-3405606"},
    {"nama": "SJN/D MOHD JURAIDI BIN SAMSUDIN", "seksyen": "E6", "no_tel": "012-6769745"},
    {"nama": "SJN/D VASTERMORAILE ANAK MONDAY", "seksyen": "E3", "no_tel": "017-7058084"},
    {"nama": "KPL/D LO YEE TYNG", "seksyen": "E4", "no_tel": "016-8374537"},
    {"nama": "KPL/D MUHAMMAD SYUKRI BIN ZULKIFLI", "seksyen": "E5", "no_tel": "012-6196573"},
    {"nama": "KPL/D YVONNE KUA CHEE WEI", "seksyen": "E3", "no_tel": "011-10905627"},
    {"nama": "KPL/D MOHAMAD PARIS BIN ROMELI", "seksyen": "E6", "no_tel": "017-3864533"},
    {"nama": "KPL/D RIZQI ALFIAN BIN KOMIN", "seksyen": "E3", "no_tel": "011-23246744"},
    {"nama": "L/KPL/D NURHIDAYAH BINTI SUHAMI", "seksyen": "E4", "no_tel": "017-7011952"},
    {"nama": "KONS NUR AMIERA BINTI MOHAMAD ALI AHAN", "seksyen": "E5", "no_tel": "017-6723075"},
    {"nama": "L/KPL AHMAD HELMI HARIRI BIN SALVEHUJRAIE", "seksyen": "E3", "no_tel": "0111-66522645"},
    {"nama": "SJN/D JAAFAR BIN MA HASAN", "seksyen": "E6", "no_tel": "011-11474013"},
    {"nama": "P/SJN/D ASRUL AZWANDI BIN AHMAD", "seksyen": "E3", "no_tel": "019-4636430"},
    {"nama": "KPL/D MOHD KHAIROL ANUAR BIN MOHAMED", "seksyen": "E4", "no_tel": "010-4358116"},
    {"nama": "KPL/D ABDUL HALIM BIN KAMARUDDIN", "seksyen": "E5", "no_tel": "013-2425850"},
    {"nama": "L/KPL/D KOLOZUM SAFINAH BINTI VARUSAI", "seksyen": "E3", "no_tel": "011-65670875"},
    {"nama": "L/KPL ATHIRAH NABILA BINTI SAEDIN", "seksyen": "E6", "no_tel": "018-3166724"},
    {"nama": "L/KPL/D HANA LYDIA BINTI LIERON", "seksyen": "E3", "no_tel": "018-8709538"},
    {"nama": "SJN/D SHAHRUL IDHAM BIN MD SALIM", "seksyen": "E4", "no_tel": "014-2440181"},
    {"nama": "KPL/D MOHD KHAIROL IZDZUAN BIN MOHAMAD ZAHARI", "seksyen": "E5", "no_tel": "012-2796394"},
    {"nama": "KPL/D SUFIAN BIN ISMAIL", "seksyen": "E3", "no_tel": "017-6217662"},
    {"nama": "KPL/D NELSON ANAK NUING", "seksyen": "E6", "no_tel": "017-3655175"},
    {"nama": "KPL/D MOHAMAD HAFIS BIN SALLEH", "seksyen": "E3", "no_tel": "013-4928247"},
    {"nama": "L/KPL/D MUHAMMAD ABDUL WAFIE HAZIQ BIN HAMDAN", "seksyen": "E4", "no_tel": "017-5431206"}
]

PENOLONG_DIKECUALIKAN_DEFAULT = [
    "KPL/D FITRI HAMIZAN BIN MOHD SHUKRI",
    "KPL/D MOHD AZHARI BIN ABD KARIM",
    "SJN/D MUHD FAIRUL BIN SIDI AHMAD",
    "KPL/D BENNY DENESIUS",
    "KPL/D MOHD TAUFIQ BIN NORHISHAM",
    "KPL/D MUHAMAD HANAFI BIN ZAINOL"
]

REKOD_SEP_2026 = {
    "INSP ROHAIZASHAFIKA BINTI MOHAMAD RADZALI": {"WD": 2, "WE": 0},
    "INSP WAN MUHAMAD MUSTAQIM BIN ABDUL RAHMAN": {"WD": 1, "WE": 1},
    "INSP MUHAMAD IQRAM BIN MOHAMAD BUKHARI": {"WD": 1, "WE": 0},
    "INSP ADDLYZAN BIN ABD MANAP": {"WD": 1, "WE": 0},
    "INSP SHURENDREN A/L JAYARAMAN": {"WD": 2, "WE": 0},
    "INSP SABARI BIN BUJANG": {"WD": 0, "WE": 1},
    "INSP MUHAMMAD AFIQ BIN NORDIN": {"WD": 1, "WE": 1},
    "INSP KHAIRUL HAFIZ BIN SHAFFIE AHMED": {"WD": 1, "WE": 1},
    "INSP DZUL FADHLI BIN ABDUL HALIM": {"WD": 0, "WE": 1},
    "INSP MOHAMAD IQBAL IBRAHIM BIN DAUD": {"WD": 1, "WE": 0},
    "INSP ABDUL MUIN BIN AB AZIZ": {"WD": 1, "WE": 1},
    "INSP AZFIZI BIN AZIZ": {"WD": 2, "WE": 0},
    "INSP AMI RUSLAN BIN MOHD RUSLAN": {"WD": 2, "WE": 0},
    "INSP FARAH NADJWA BT MOHD LASA": {"WD": 0, "WE": 2},
    "INSP MOHD JEFFERE BIN ALI": {"WD": 1, "WE": 1},
    "INSP NUR AADILA BINTI ABDUL MALEK": {"WD": 0, "WE": 1},
    "INSP EZYANTI BT MUHAMAD": {"WD": 1, "WE": 1},
    "INSP MOHD FAIZ BIN ESA": {"WD": 2, "WE": 0},
    "INSP KASHMINDEJRIT KAUR A/P CHARLES": {"WD": 0, "WE": 1},
    "INSP FAIZ BIN BASAR": {"WD": 0, "WE": 1},
    "INSP MOHAMAD SYARUL FADLI BIN MAT HUSAIN": {"WD": 1, "WE": 0},
    "INSP NOR HAIZAH BINTI YACOB": {"WD": 0, "WE": 1},
    "INSP MOHD RAZIF BIN ROSSMAN": {"WD": 0, "WE": 1},
    "INSP NORAINUN MUBIN BT ANUAR": {"WD": 0, "WE": 1},
    "INSP KUKENDRAN A/L YOGENDRAN": {"WD": 0, "WE": 1},
    "INSP MOHD SHAHRIL BIN MOHAMED RESALI": {"WD": 0, "WE": 1},
    "INSP MOHD SHAHRILBUN BIN AMDAN": {"WD": 1, "WE": 0},
    "INSP MOHD RUSDI BIN HASSAN": {"WD": 1, "WE": 0},
    "INSP MUHAMAD ARIFF BIN ABD RAHIM": {"WD": 0, "WE": 0},
    "INSP MUHAMMAD NOR QHAIREIL BIN NOR AZMI": {"WD": 0, "WE": 0},
    "INSP NORHAZIERAM BINTI ZAKARIA": {"WD": 0, "WE": 0}
}

# ==========================================
# 3. ANTARAMUKA STREAMLIT
# ==========================================
tab1, tab2 = st.tabs(["⚙️ Penjanaan Jadual Baru", "📜 Rekod Sejarah & Format Cetak"])

with tab1:
    col1, col2 = st.columns(2)
    with col1:
        bulan = st.selectbox("Pilih Bulan", list(range(1, 13)), index=9)
    with col2:
        tahun = st.number_input("Tahun", min_value=2024, max_value=2030, value=datetime.datetime.now().year)

    num_days = calendar.monthrange(tahun, bulan)[1]
    senarai_tarikh = [datetime.date(tahun, bulan, d) for d in range(1, num_days + 1)]

    st.sidebar.header("🔄 Tetapan Giliran")
    aktif_giliran_lelaki = st.sidebar.checkbox("Aktifkan Giliran WE/WD Pegawai Lelaki", value=True)

    st.sidebar.header("🚫 Pengecualian Tugas Pegawai (Insp)")
    pegawai_dikecualikan = st.sidebar.multiselect("Pegawai Dikecualikan:", [p["nama"] for p in PEGAWAI])

    st.sidebar.header("🚫 Pengecualian Tugas Penolong Pegawai")
    penolong_dikecualikan = st.sidebar.multiselect(
        "Penolong Pegawai Dikecualikan:",
        [p["nama"] for p in PENOLONG_PEGAWAI],
        default=[p for p in PENOLONG_DIKECUALIKAN_DEFAULT if any(x["nama"] == p for x in PENOLONG_PEGAWAI)]
    )

    st.sidebar.header("🎉 Tetapan Cuti Umum")
    tarikh_cuti_umum = st.sidebar.multiselect(
        "Pilih tarikh Cuti Umum:",
        options=senarai_tarikh,
        format_func=lambda d: d.strftime("%d/%m/%Y (%A)")
    )

    st.sidebar.header("📌 Special Assignment (Manual Override)")
    override_pegawai = {}
    override_penolong = {}
    
    senarai_pegawai_aktif = [p["nama"] for p in PEGAWAI if p["nama"] not in pegawai_dikecualikan]
    senarai_penolong_aktif = [p["nama"] for p in PENOLONG_PEGAWAI if p["nama"] not in penolong_dikecualikan]

    with st.sidebar.expander("Tetapan Manual Override (Jika Ada)"):
        for dt in senarai_tarikh:
            is_weekend = dt.weekday() in [5, 6] or dt in tarikh_cuti_umum
            tkh_str = dt.strftime("%d/%m/%Y")
            
            if is_weekend:
                s_peg = st.selectbox(f"{tkh_str} (Siang) - Pegawai", ["- Auto -"] + senarai_pegawai_aktif, key=f"peg_{tkh_str}_S")
                if s_peg != "- Auto -": override_pegawai[f"{tkh_str}_Siang"] = s_peg
                
                s_pen = st.selectbox(f"{tkh_str} (Siang) - Penolong", ["- Auto -"] + senarai_penolong_aktif, key=f"pen_{tkh_str}_S")
                if s_pen != "- Auto -": override_penolong[f"{tkh_str}_Siang"] = s_pen

                m_peg = st.selectbox(f"{tkh_str} (Malam) - Pegawai", ["- Auto -"] + senarai_pegawai_aktif, key=f"peg_{tkh_str}_M")
                if m_peg != "- Auto -": override_pegawai[f"{tkh_str}_Malam"] = m_peg
                
                m_pen = st.selectbox(f"{tkh_str} (Malam) - Penolong", ["- Auto -"] + senarai_penolong_aktif, key=f"pen_{tkh_str}_M")
                if m_pen != "- Auto -": override_penolong[f"{tkh_str}_Malam"] = m_pen
            else:
                b_peg = st.selectbox(f"{tkh_str} (Biasa) - Pegawai", ["- Auto -"] + senarai_pegawai_aktif, key=f"peg_{tkh_str}_B")
                if b_peg != "- Auto -": override_pegawai[f"{tkh_str}_Biasa"] = b_peg
                
                b_pen = st.selectbox(f"{tkh_str} (Biasa) - Penolong", ["- Auto -"] + senarai_penolong_aktif, key=f"pen_{tkh_str}_B")
                if b_pen != "- Auto -": override_penolong[f"{tkh_str}_Biasa"] = b_pen

    if st.button("🚀 Jana Jadual Penugasan", type="primary"):
        tugas_wd_peg = {p["nama"]: 0 for p in PEGAWAI if p["nama"] not in pegawai_dikecualikan}
        tugas_we_peg = {p["nama"]: 0 for p in PEGAWAI if p["nama"] not in pegawai_dikecualikan}
        
        tugas_penolong = {p["nama"]: 0 for p in PENOLONG_PEGAWAI if p["nama"] not in penolong_dikecualikan}

        pegawai_tugas_jumaat = set()
        pegawai_hari_sebelumnya = set()
        penolong_hari_sebelumnya = set()
        jadual = []

        bulan_lepas = 12 if bulan == 1 else bulan - 1
        tahun_lepas = tahun - 1 if bulan == 1 else tahun
        
        rekod_bulan_lepas = REKOD_SEP_2026 if (bulan_lepas == 9 and tahun_lepas == 2026) else {}

        def pilih_pegawai(is_wknd, elak_set, is_jumaat=False):
            calon = [p for p in senarai_pegawai_aktif if p not in elak_set]

            def dapatkan_had_maksimum(nama_pegawai):
                rec_lepas = rekod_bulan_lepas.get(nama_pegawai, {"WD": 0, "WE": 0})
                jumlah_lepas = rec_lepas.get("WD", 0) + rec_lepas.get("WE", 0)
                return 1 if jumlah_lepas >= 2 else 2

            calon_cap = [p for p in calon if (tugas_wd_peg.get(p, 0) + tugas_we_peg.get(p, 0)) < dapatkan_had_maksimum(p)]
            if calon_cap: calon = calon_cap

            if bulan == 10 and "INSP MOHD FAIZ BIN ESA" in calon:
                if is_wknd or tugas_wd_peg.get("INSP MOHD FAIZ BIN ESA", 0) >= 1:
                    if len(calon) > 1: calon = [p for p in calon if p != "INSP MOHD FAIZ BIN ESA"]

            if aktif_giliran_lelaki:
                if is_wknd:
                    calon_we = []
                    for p in calon:
                        rec_we_lepas = rekod_bulan_lepas.get(p, {}).get("WE", 0)
                        sudah_we = tugas_we_peg.get(p, 0)
                        if is_pegawai_wanita(p):
                            if sudah_we < 1: calon_we.append(p)
                        else:
                            if rec_we_lepas == 0 and sudah_we < 1: calon_we.append(p)
                    if calon_we: calon = calon_we
                else:
                    calon_wd = [p for p in calon if not is_pegawai_wanita(p) and rekod_bulan_lepas.get(p, {}).get("WE", 0) > 0]
                    if len(calon_wd) >= 1: calon = calon_wd

            if is_wknd:
                calon_no_jumaat = [p for p in calon if p not in pegawai_tugas_jumaat]
                if calon_no_jumaat: calon = calon_no_jumaat

            if not calon: calon = [p for p in senarai_pegawai_aktif if p not in elak_set] or senarai_pegawai_aktif

            def dapatkan_skor(nama_pegawai):
                rec = rekod_bulan_lepas.get(nama_pegawai, {"WD": 0, "WE": 0})
                if is_wknd: return (tugas_we_peg[nama_pegawai] * 100) + (rec["WE"] * 10) + (tugas_wd_peg[nama_pegawai] + tugas_we_peg[nama_pegawai])
                else: return (tugas_wd_peg[nama_pegawai] * 100) + (rec["WD"] * 10) + (tugas_wd_peg[nama_pegawai] + tugas_we_peg[nama_pegawai])

            min_skor = min([dapatkan_skor(c) for c in calon])
            final = [c for c in calon if dapatkan_skor(c) == min_skor]
            pilihan = random.choice(final)

            if is_wknd: tugas_we_peg[pilihan] += 1
            else: tugas_wd_peg[pilihan] += 1
            if is_jumaat: pegawai_tugas_jumaat.add(pilihan)
            
            return pilihan

        def pilih_penolong(elak_set):
            calon = [p for p in senarai_penolong_aktif if p not in elak_set]
            calon_1kali = [p for p in calon if tugas_penolong.get(p, 0) < 1]
            if calon_1kali: calon = calon_1kali
            if not calon: calon = [p for p in senarai_penolong_aktif if p not in elak_set] or senarai_penolong_aktif

            min_tugas = min([tugas_penolong[c] for c in calon])
            final = [c for c in calon if tugas_penolong[c] == min_tugas]
            pilihan = random.choice(final)
            tugas_penolong[pilihan] += 1
            return pilihan

        for dt in senarai_tarikh:
            is_weekend = dt.weekday() in [5, 6] or dt in tarikh_cuti_umum
            is_jumaat = (dt.weekday() == 4 and dt not in tarikh_cuti_umum)
            tkh_str = dt.strftime("%d/%m/%Y")
            hari_str = dt.strftime("%A") + (" (CUTI UMUM)" if dt in tarikh_cuti_umum else "")
            
            pegawai_hari_ini = set()
            penolong_hari_ini = set()

            if is_weekend:
                # Syif Siang
                p_s = override_pegawai.get(f"{tkh_str}_Siang") or pilih_pegawai(True, pegawai_hari_sebelumnya)
                if f"{tkh_str}_Siang" in override_pegawai: tugas_we_peg[p_s] += 1
                pegawai_hari_ini.add(p_s)

                pen_s = override_penolong.get(f"{tkh_str}_Siang") or pilih_penolong(penolong_hari_sebelumnya)
                if f"{tkh_str}_Siang" in override_penolong: tugas_penolong[pen_s] += 1
                penolong_hari_ini.add(pen_s)

                tel_p_s = next(p["no_tel"] for p in PEGAWAI if p["nama"] == p_s)
                sek_p_s = next(p["seksyen"] for p in PEGAWAI if p["nama"] == p_s)
                
                tel_pen_s = next(p["no_tel"] for p in PENOLONG_PEGAWAI if p["nama"] == pen_s)
                sek_pen_s = next(p["seksyen"] for p in PENOLONG_PEGAWAI if p["nama"] == pen_s)

                jadual.append({
                    "Tarikh": tkh_str, "Hari": hari_str, "Syif / Masa": "Siang (0800 - 2000)",
                    "Nama Pegawai": p_s, "Seksyen Pegawai": sek_p_s, "No. Tel Pegawai": tel_p_s,
                    "Nama Penolong": pen_s, "Seksyen Penolong": sek_pen_s, "No. Tel Penolong": tel_pen_s
                })

                # Syif Malam
                p_m = override_pegawai.get(f"{tkh_str}_Malam") or pilih_pegawai(True, pegawai_hari_sebelumnya.union(pegawai_hari_ini))
                if f"{tkh_str}_Malam" in override_pegawai: tugas_we_peg[p_m] += 1
                pegawai_hari_ini.add(p_m)

                pen_m = override_penolong.get(f"{tkh_str}_Malam") or pilih_penolong(penolong_hari_sebelumnya.union(penolong_hari_ini))
                if f"{tkh_str}_Malam" in override_penolong: tugas_penolong[pen_m] += 1
                penolong_hari_ini.add(pen_m)

                tel_p_m = next(p["no_tel"] for p in PEGAWAI if p["nama"] == p_m)
                sek_p_m = next(p["seksyen"] for p in PEGAWAI if p["nama"] == p_m)
                
                tel_pen_m = next(p["no_tel"] for p in PENOLONG_PEGAWAI if p["nama"] == pen_m)
                sek_pen_m = next(p["seksyen"] for p in PENOLONG_PEGAWAI if p["nama"] == pen_m)

                jadual.append({
                    "Tarikh": tkh_str, "Hari": hari_str, "Syif / Masa": "Malam (2000 - 0800)",
                    "Nama Pegawai": p_m, "Seksyen Pegawai": sek_p_m, "No. Tel Pegawai": tel_p_m,
                    "Nama Penolong": pen_m, "Seksyen Penolong": sek_pen_m, "No. Tel Penolong": tel_pen_m
                })

            else:
                # Syif Biasa
                p_b = override_pegawai.get(f"{tkh_str}_Biasa") or pilih_pegawai(False, pegawai_hari_sebelumnya, is_jumaat=is_jumaat)
                if f"{tkh_str}_Biasa" in override_pegawai: 
                    tugas_wd_peg[p_b] += 1
                    if is_jumaat: pegawai_tugas_jumaat.add(p_b)
                pegawai_hari_ini.add(p_b)

                pen_b = override_penolong.get(f"{tkh_str}_Biasa") or pilih_penolong(penolong_hari_sebelumnya)
                if f"{tkh_str}_Biasa" in override_penolong: tugas_penolong[pen_b] += 1
                penolong_hari_ini.add(pen_b)

                tel_p_b = next(p["no_tel"] for p in PEGAWAI if p["nama"] == p_b)
                sek_p_b = next(p["seksyen"] for p in PEGAWAI if p["nama"] == p_b)
                
                tel_pen_b = next(p["no_tel"] for p in PENOLONG_PEGAWAI if p["nama"] == pen_b)
                sek_pen_b = next(p["seksyen"] for p in PENOLONG_PEGAWAI if p["nama"] == pen_b)

                jadual.append({
                    "Tarikh": tkh_str, "Hari": hari_str, "Syif / Masa": "Biasa (1700 - 0800)",
                    "Nama Pegawai": p_b, "Seksyen Pegawai": sek_p_b, "No. Tel Pegawai": tel_p_b,
                    "Nama Penolong": pen_b, "Seksyen Penolong": sek_pen_b, "No. Tel Penolong": tel_pen_b
                })

            pegawai_hari_sebelumnya = pegawai_hari_ini
            penolong_hari_sebelumnya = penolong_hari_ini

        st.session_state['df_jadual_temp'] = pd.DataFrame(jadual)
        st.session_state['bulan_temp'] = bulan
        st.session_state['tahun_temp'] = tahun

    if 'df_jadual_temp' in st.session_state and st.session_state['bulan_temp'] == bulan and st.session_state['tahun_temp'] == tahun:
        df_j = st.session_state['df_jadual_temp']
        st.subheader(f"📅 Hasil Jadual Penugasan (Bulan {bulan}/{tahun})")
        st.dataframe(df_j, use_container_width=True)

        if st.button("🔒 Sahkan & Simpan Jadual Ini", type="secondary"):
            simpan_jadual_db(bulan, tahun, df_j)
            st.success(f"✅ Jadual bagi Bulan {bulan}/{tahun} BERHASIL DISIMPAN!")

with tab2:
    st.subheader("📜 Carian & Paparan Format Cetak (Lampiran 'A')")
    col_a, col_b = st.columns(2)
    with col_a:
        bulan_cari = st.selectbox("Pilih Bulan Rekod", list(range(1, 13)), index=8, key="b_cari")
    with col_b:
        tahun_cari = st.number_input("Pilih Tahun Rekod", min_value=2024, max_value=2030, value=2026, key="t_cari")

    df_rekod = ambil_jadual_db(bulan_cari, tahun_cari)
    
    if not df_rekod.empty:
        st.info(f"Rekod Ditemui. Tarikh Disahkan: {df_rekod['Tarikh Disahkan'].iloc[0]}")
        html_code = jana_html_lampiran_a(df_rekod, bulan_cari, tahun_cari)
        word_data = jana_word_lampiran_a(df_rekod, bulan_cari, tahun_cari)
        
        col_dl1, col_dl2, col_dl3, col_del = st.columns(4)
        with col_dl1:
            st.download_button("📄 Muat Turun (Word)", data=word_data, file_name=f"Lampiran_A_SBCC_{bulan_cari}_{tahun_cari}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        with col_dl2:
            st.download_button("🖨️ Muat Turun (HTML)", html_code, f"Lampiran_A_{bulan_cari}_{tahun_cari}.html", "text/html")
        with col_dl3:
            st.download_button("📥 Muat Turun (CSV)", df_rekod.to_csv(index=False).encode('utf-8'), f"Jadual_{bulan_cari}_{tahun_cari}.csv", "text/csv")
        with col_del:
            if st.button("🗑️ Padam Rekod Ini", type="primary"):
                padam_jadual_db(bulan_cari, tahun_cari)
                st.success("Rekod dipadam.")
                st.rerun()

        st.markdown("---")
        st.subheader("📄 Pratonton Dokumen Cetakan (Format Lampiran 'A'):")
        st.components.v1.html(html_code, height=600, scrolling=True)
    else:
        st.warning(f"Tiada rekod jadual yang disahkan ditemui bagi Bulan {bulan_cari}/{tahun_cari}.")
